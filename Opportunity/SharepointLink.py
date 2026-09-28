import os
import re
import urllib.parse
import openpyxl


def parse_replacement_input(user_input_str):
    """
    Parses user replacement rules from input text.
    Supports formats like:
      "01-Denver/02-Opportunities"="DEN/Oppy"
      "old1"="new1", "old2"="new2"
      "old1", "new1"
    """
    replacements = {}
    if not user_input_str or not user_input_str.strip():
        return replacements

    quoted_tokens = re.findall(r'"([^"]*)"', user_input_str)

    if len(quoted_tokens) >= 2 and len(
            quoted_tokens) % 2 == 0 and '=' not in user_input_str and ':' not in user_input_str:
        for i in range(0, len(quoted_tokens), 2):
            replacements[quoted_tokens[i]] = quoted_tokens[i + 1]
    else:
        pairs = user_input_str.split(',')
        for pair in pairs:
            if '=' in pair:
                parts = pair.split('=', 1)
            elif ':' in pair:
                parts = pair.split(':', 1)
            else:
                continue
            old_val = parts[0].strip().strip('"\'')
            new_val = parts[1].strip().strip('"\'')
            if old_val:
                replacements[old_val] = new_val

    return replacements


def process_excel(file_path, replacements=None, output_file=None):
    """
    Processes all rows in the Excel file:
      1. Applies URL string replacements to Sharepoint Link.
      2. Removes prefix and suffix from the folder segment in Updated Sharepoint Link so only the Opportunity Number remains.
      3. Keeps only the Opportunity Number in 'Kept Opportunity Number'.
      4. Stores the removed prefix/suffix text in 'Removed Portion'.
    """
    if not os.path.exists(file_path):
        print(f"\n[Error] File not found at '{file_path}'")
        return None

    if replacements is None:
        replacements = {}
    elif isinstance(replacements, str):
        replacements = parse_replacement_input(replacements)

    print(f"\nActive Replacement Rules: {replacements if replacements else 'None (No replacements)'}")

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    # Identify header columns
    header = [cell.value for cell in ws[1]]
    opty_col_idx = header.index("Opportunity Number") + 1 if "Opportunity Number" in header else 1
    link_col_idx = header.index("Sharepoint Link") + 1 if "Sharepoint Link" in header else 2

    # Add output headers
    new_headers = ["Updated Sharepoint Link", "Kept Opportunity Number", "Removed Portion", "Status"]
    header_vals = [cell.value for cell in ws[1]]

    if "Kept Opportunity Number" not in header_vals:
        start_new_col = ws.max_column + 1
        for idx, h_text in enumerate(new_headers):
            ws.cell(row=1, column=start_new_col + idx, value=h_text)
        col_link_out = start_new_col
        col_opty_out = start_new_col + 1
        col_rem_out = start_new_col + 2
        col_stat_out = start_new_col + 3
    else:
        col_link_out = header_vals.index("Updated Sharepoint Link") + 1
        col_opty_out = header_vals.index("Kept Opportunity Number") + 1
        col_rem_out = header_vals.index("Removed Portion") + 1
        col_stat_out = header_vals.index("Status") + 1

    processed_count = 0
    matched_count = 0

    print("\n--- Processing All Rows ---")
    for row in range(2, ws.max_row + 1):
        opty_val = ws.cell(row=row, column=opty_col_idx).value
        link_val = ws.cell(row=row, column=link_col_idx).value

        if opty_val is None and link_val is None:
            continue

        opty_str = str(opty_val).strip() if opty_val is not None else ""
        link_str = str(link_val).strip() if link_val is not None else ""

        # Step 1: Apply URL text replacements
        updated_link = link_str
        for old_str, new_str in replacements.items():
            if old_str:
                updated_link = updated_link.replace(old_str, new_str)

        # Step 2: Extract folder segment and decode URL %20 encoding
        folder_segment = updated_link.split('/')[-1] if updated_link else ""
        decoded_folder = urllib.parse.unquote(folder_segment)

        # Step 3: Check Opportunity Number match & clean Updated Sharepoint Link
        if opty_str and opty_str in decoded_folder:
            kept_opty = opty_str
            # Removed portion stores prefix/suffix removed from folder name
            removed_portion = decoded_folder.replace(opty_str, '', 1)

            # Build Updated Sharepoint Link with prefix/suffix REMOVED (ending in only the Opportunity Number)
            base_url = updated_link[:updated_link.rfind('/') + 1]
            final_updated_link = base_url + opty_str
            status = "Matched"
            matched_count += 1
        elif opty_str and opty_str in updated_link:
            kept_opty = opty_str
            removed_portion = updated_link.replace(opty_str, '', 1)
            final_updated_link = updated_link
            status = "Matched in URL"
            matched_count += 1
        else:
            kept_opty = "Not Found"
            removed_portion = ""
            final_updated_link = updated_link
            status = "Not Found"

        # Write results into output columns
        ws.cell(row=row, column=col_link_out, value=final_updated_link)
        ws.cell(row=row, column=col_opty_out, value=kept_opty)
        ws.cell(row=row, column=col_rem_out, value=removed_portion)
        ws.cell(row=row, column=col_stat_out, value=status)

        processed_count += 1
        if row <= 6 or row == ws.max_row:
            print(
                f"Row {row:2d} | Original Opty: {opty_str:6s} | Kept: {kept_opty:9s} | Updated Link: .../{final_updated_link.split('/')[-1]} | Removed: {removed_portion[:30]:30s}")

    if output_file is None:
        base, ext = os.path.splitext(file_path)
        output_file = f"{base}_Processed{ext}"

    try:
        wb.save(output_file)
    except PermissionError:
        base, ext = os.path.splitext(output_file)
        output_file = f"{base}_v2{ext}"
        wb.save(output_file)

    print(f"\nProcessing Complete!")
    print(f"Total Rows Processed: {processed_count}")
    print(f"Opportunity Matches: {matched_count}")
    print(f"Output saved to: {output_file}")
    return output_file


if __name__ == "__main__":
    print("=" * 60)
    print("      Excel Sharepoint Link & Opportunity Processor      ")
    print("=" * 60)

    default_file = r"C:\Users\vishal walunj\Downloads\EBS_PROD_OPPTY.xlsx"

    print(f"\nDefault Excel file: {default_file}")
    file_input = input("Enter Excel file path (or press ENTER to use default): ").strip().strip('"\'')

    stray_find_text = ""
    if not file_input:
        file_path = default_file
    elif os.path.exists(file_input):
        file_path = file_input
    else:
        print(f"\n[Note] Path '{file_input}' not found as a file.")
        if os.path.exists(default_file):
            print(f"-> Using default Excel file: {default_file}")
            file_path = default_file
            stray_find_text = file_input
        else:
            print("[Error] Default Excel file not found either. Please check file path.")
            exit(1)

    replacements = {}

    if stray_find_text:
        print(f"\nDetected Find Text: '{stray_find_text}'")
        rep_val = input(f"Enter replacement text for '{stray_find_text}' (or press ENTER to skip): ").strip().strip(
            '"\'')
        if rep_val:
            replacements[stray_find_text] = rep_val
            print(f"-> Added replacement rule: '{stray_find_text}' -> '{rep_val}'")

    print("\n--- URL Replacements ---")
    print("Leave 'Find Text' blank and press ENTER when done / to skip replacements.")

    while True:
        find_text = input("\nEnter text to FIND in URL (or press ENTER to finish): ").strip().strip('"\'')
        if not find_text:
            break

        parsed = parse_replacement_input(find_text)
        if parsed:
            replacements.update(parsed)
            for k, v in parsed.items():
                print(f"-> Added replacement rule: '{k}' -> '{v}'")
            continue

        replace_text = input(f"Enter text to REPLACE '{find_text}' with: ").strip().strip('"\'')
        replacements[find_text] = replace_text
        print(f"-> Added replacement rule: '{find_text}' -> '{replace_text}'")

    process_excel(file_path, replacements)