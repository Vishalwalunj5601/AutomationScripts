import pandas as pd
import re
from itertools import combinations
from rapidfuzz import fuzz
from openpyxl import load_workbook
from openpyxl.styles import Font

# ==========================================================
# CONFIGURATION
# ==========================================================

INPUT_FILE = r"C:\Users\Vishal Walunj\Downloads\DuplicatePartySupplier.xlsx"
OUTPUT_FILE = r"C:\AIR_CONTROL_CONCEPTS\DEV3-PROD_1C\Report\DuplicateParty\Customer_Duplicate_Report_su.xlsx"

PARTY_THRESHOLD = 80
ADDRESS_THRESHOLD = 95

# ==========================================================
# HELPER FUNCTIONS
# ==========================================================

def normalize_text(text):
    if pd.isna(text):
        return ""

    text = str(text).upper().strip()

    text = re.sub(r'[^A-Z0-9\s]', ' ', text)
    text = re.sub(r'\s+', ' ', text)

    return text


def create_full_address(row):
    parts = [
        row.get("ADDRESS1", ""),
        row.get("ADDRESS2", ""),
        row.get("ADDRESS3", ""),
        row.get("ADDRESS4", ""),
        row.get("CITY", ""),
        row.get("STATE", ""),
        row.get("POSTAL_CODE", ""),
        row.get("COUNTRY", "")
    ]

    return normalize_text(
        " ".join([str(x) for x in parts if pd.notna(x)])
    )


def get_match_category(party_score, address_score):
    if party_score == 100 and address_score == 100:
        return "EXACT PARTY + EXACT ADDRESS"

    elif party_score == 100 and address_score >= 95:
        return "EXACT PARTY + FUZZY ADDRESS"

    elif party_score >= 80 and address_score == 100:
        return "FUZZY PARTY + EXACT ADDRESS"

    elif party_score >= 80 and address_score >= 95:
        return "FUZZY PARTY + FUZZY ADDRESS"

    return "NO MATCH"


# ==========================================================
# READ FILE
# ==========================================================

print("Reading input file...")

df = pd.read_excel(INPUT_FILE)

df.columns = [c.strip().upper() for c in df.columns]

required_cols = [
    "OPCO",
    "PARTY_NAME",
    "PARTY_NUMBER",
    "ADDRESS1",
    "ADDRESS2",
    "ADDRESS3",
    "ADDRESS4",
    "CITY",
    "POSTAL_CODE",
    "STATE",
    "COUNTRY",
    "SITE_NAME"
]

missing_cols = [c for c in required_cols if c not in df.columns]

if missing_cols:
    raise Exception(f"Missing columns: {missing_cols}")

# ==========================================================
# NORMALIZE DATA
# ==========================================================

print("Normalizing data...")

df["PARTY_NAME_NORM"] = df["PARTY_NAME"].apply(normalize_text)
df["FULL_ADDRESS"] = df.apply(create_full_address, axis=1)

# ==========================================================
# SHEET 1 : EXACT PARTY MATCHES (100%)
# ==========================================================

print("Building Exact Party Matches...")

exact_party_matches = []

party_groups = df.groupby(["PARTY_NUMBER", "PARTY_NAME_NORM"])

for (party_number, party_name), group in party_groups:

    opcos = group["OPCO"].astype(str).unique()

    if len(opcos) > 1:

        exact_party_matches.append({
            "PARTY_NUMBER": party_number,
            "PARTY_NAME": party_name,
            "OPCOS": ", ".join(sorted(opcos)),
            "OPCO_COUNT": len(opcos),
            "PARTY_SCORE": 100
        })

exact_party_df = pd.DataFrame(exact_party_matches)

# ==========================================================
# SHEET 2 : FUZZY PARTY MATCHES (80-99%)
# ==========================================================

print("Building Fuzzy Party Matches...")

fuzzy_party_matches = []

party_master = (
    df[
        ["PARTY_NAME",
         "PARTY_NAME_NORM",
         "PARTY_NUMBER",
         "OPCO"]
    ]
    .drop_duplicates()
    .reset_index(drop=True)
)

for i in range(len(party_master)):

    row1 = party_master.iloc[i]

    for j in range(i + 1, len(party_master)):

        row2 = party_master.iloc[j]

        if row1["OPCO"] == row2["OPCO"]:
            continue

        score = fuzz.token_sort_ratio(
            row1["PARTY_NAME_NORM"],
            row2["PARTY_NAME_NORM"]
        )

        if PARTY_THRESHOLD <= score < 100:

            fuzzy_party_matches.append({
                "PARTY_NAME_1": row1["PARTY_NAME"],
                "PARTY_NUMBER_1": row1["PARTY_NUMBER"],
                "OPCO_1": row1["OPCO"],
                "PARTY_NAME_2": row2["PARTY_NAME"],
                "PARTY_NUMBER_2": row2["PARTY_NUMBER"],
                "OPCO_2": row2["OPCO"],
                "PARTY_SCORE": score
            })

fuzzy_party_df = pd.DataFrame(fuzzy_party_matches)

# ==========================================================
# SHEET 3 : PARTY + ADDRESS MATCHES
# ==========================================================

print("Building Party + Address Matches...")

party_address_matches = []

for (party_number, party_name), group in party_groups:

    opcos = group["OPCO"].unique()

    if len(opcos) < 2:
        continue

    records = group.to_dict("records")

    for r1, r2 in combinations(records, 2):

        if r1["OPCO"] == r2["OPCO"]:
            continue

        address_score = fuzz.token_sort_ratio(
            r1["FULL_ADDRESS"],
            r2["FULL_ADDRESS"]
        )

        if address_score >= ADDRESS_THRESHOLD:

            party_score = 100

            party_address_matches.append({

                "PARTY_NUMBER": party_number,
                "PARTY_NAME": party_name,

                "OPCO_1": r1["OPCO"],
                "OPCO_2": r2["OPCO"],

                "SITE_1": r1["SITE_NAME"],
                "SITE_2": r2["SITE_NAME"],

                "PARTY_SCORE": party_score,
                "ADDRESS_SCORE": address_score,

                "MATCH_CATEGORY":
                    get_match_category(
                        party_score,
                        address_score
                    ),

                "ADDRESS_1": r1["FULL_ADDRESS"],
                "ADDRESS_2": r2["FULL_ADDRESS"]
            })

party_address_df = pd.DataFrame(
    party_address_matches,
    columns=[
        "PARTY_NUMBER",
        "PARTY_NAME",
        "OPCO_1",
        "OPCO_2",
        "SITE_1",
        "SITE_2",
        "PARTY_SCORE",
        "ADDRESS_SCORE",
        "MATCH_CATEGORY",
        "ADDRESS_1",
        "ADDRESS_2"
    ]
)

# ==========================================================
# SHEET 4 : OPCO SUMMARY
# ==========================================================

print("Building OPCO Summary...")

opco_summary = []

duplicate_party_numbers = set()

for _, row in exact_party_df.iterrows():
    duplicate_party_numbers.add(row["PARTY_NUMBER"])

for opco in sorted(df["OPCO"].astype(str).unique()):

    opco_df = df[df["OPCO"].astype(str) == opco]

    total_records = len(opco_df)

    total_parties = opco_df["PARTY_NUMBER"].nunique()

    duplicate_parties = (
        opco_df[
            opco_df["PARTY_NUMBER"]
            .isin(duplicate_party_numbers)
        ]["PARTY_NUMBER"]
        .nunique()
    )

    if party_address_df.empty:
        duplicate_addresses = 0
    else:
        duplicate_addresses = len(
            party_address_df[
                (party_address_df["OPCO_1"] == opco) |
                (party_address_df["OPCO_2"] == opco)
                ]
        )

    duplicate_pct = round(
        (
            duplicate_parties / total_parties
        ) * 100,
        2
    ) if total_parties > 0 else 0

    opco_summary.append({

        "OPCO": opco,
        "TOTAL_RECORDS": total_records,
        "TOTAL_PARTIES": total_parties,

        "DUPLICATE_PARTIES":
            duplicate_parties,

        "DUPLICATE_ADDRESSES":
            duplicate_addresses,

        "DUPLICATE_PERCENT":
            duplicate_pct
    })

opco_summary_df = pd.DataFrame(opco_summary)

# ==========================================================
# SHEET 5 : EXECUTIVE SUMMARY
# ==========================================================

summary_df = pd.DataFrame([

    {
        "METRIC": "Total Records",
        "VALUE": len(df)
    },

    {
        "METRIC": "Unique Parties",
        "VALUE": df["PARTY_NUMBER"].nunique()
    },

    {
        "METRIC": "Exact Party Matches",
        "VALUE": len(exact_party_df)
    },

    {
        "METRIC": "Fuzzy Party Matches",
        "VALUE": len(fuzzy_party_df)
    },

    {
        "METRIC": "Party + Address Matches",
        "VALUE": len(party_address_df)
    },

    {
        "METRIC": "OPCO Count",
        "VALUE": df["OPCO"].nunique()
    }

])

# ==========================================================
# WRITE EXCEL
# ==========================================================

print("Writing report...")

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    exact_party_df.to_excel(
        writer,
        sheet_name="Exact Party Match",
        index=False
    )

    fuzzy_party_df.to_excel(
        writer,
        sheet_name="Fuzzy Party Match",
        index=False
    )

    party_address_df.to_excel(
        writer,
        sheet_name="Party_Address_Match",
        index=False
    )

    opco_summary_df.to_excel(
        writer,
        sheet_name="OPCO Summary",
        index=False
    )

    summary_df.to_excel(
        writer,
        sheet_name="Executive Summary",
        index=False
    )

# ==========================================================
# FORMAT EXCEL
# ==========================================================

wb = load_workbook(OUTPUT_FILE)

for ws in wb.worksheets:

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for column in ws.columns:

        max_len = 0

        column_letter = column[0].column_letter

        for cell in column:

            try:
                if cell.value:
                    max_len = max(
                        max_len,
                        len(str(cell.value))
                    )
            except:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(max_len + 3, 60)

wb.save(OUTPUT_FILE)

print("=" * 60)
print("DUPLICATE ANALYSIS REPORT GENERATED")
print(f"Output File : {OUTPUT_FILE}")
print("=" * 60)