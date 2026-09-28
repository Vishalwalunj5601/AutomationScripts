"""
===============================================================================
Airetech PPM Conversion — Billing Events FBDI Generation Script
===============================================================================
Target Directory:
  C:\\Users\\vishal walunj\\Vishal\\1D_AIR_CONTROL_CONCEPTS\\AiretechProjects\\FINAL LOAD\\DATAFILE\\BillingEvents Import

Source File:
  C:\\Users\\vishal walunj\\Vishal\\1D_AIR_CONTROL_CONCEPTS\\AiretechProjects\\FINAL LOAD\\SOURCE\\TAB Pull from Jonas 9-18.xlsx

Outputs Generated:
  1. Paid Billing Events FBDI Workbook (.xlsx):
     - PaidBillingEvents_Populated.xlsx
  2. Unpaid Billing Events FBDI Workbook (.xlsx):
     - UnpaidBillingEvents_Populated.xlsx
  3. Macro-Enabled FBDI Workbooks (.xlsm):
     - PaidCreateBillingEventsTemplate_FINAL_LOAD.xlsm
     - UnPaidCreateBillingEventsTemplate_FINAL_LOAD.xlsm
  4. Interface CSV Files:
     - Paid_PjbBillingEventsXface.csv
     - Unpaid_PjbBillingEventsXface.csv

Key Rules:
  - Paid Event Amount = Rev Billed - A/R Owing (Filtered for Paid Amount > 0)
  - Unpaid Event Amount = A/R Owing (Filtered for Unpaid Amount > 0)
  - Event Type: "AIR Conversion"
  - Source Name: "Data Conversion"
  - Organization: "Airetech"
  - Contract Type: "Airetech_Customer Contract for Project"
  - Contract Line Number: "Conversion Contract Line"
  - Completion Date: "07/31/2026" (MM/DD/YYYY cutover date)
  - Currency: "USD"
  - Standalone Execution: Runs cleanly in PyCharm without requiring external mapping folders.
===============================================================================
"""

import os
import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# =============================================================================
# 1. DIRECTORIES & FILE PATH CONFIGURATION
# =============================================================================
OUTPUT_DATA_DIR = r"C:\Users\vishal walunj\Vishal\1D_AIR_CONTROL_CONCEPTS\AiretechProjects\FINAL LOAD\DATAFILE\BillingEvents Import"
os.makedirs(OUTPUT_DATA_DIR, exist_ok=True)

# Candidate Source Files
CANDIDATE_SOURCES = [
    r"C:\Users\vishal walunj\Vishal\1D_AIR_CONTROL_CONCEPTS\AiretechProjects\FINAL LOAD\SOURCE\TAB Pull from Jonas 9-18.xlsx",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "TAB Pull from Jonas 9-18.xlsx"),
    r"C:\Users\vishal walunj\Vishal\1D_AIR_CONTROL_CONCEPTS\AiretechProjects\DEV1\SourceFile\Airetech TAB jobs as of 9-11-26.xlsx",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "Airetech TAB jobs as of 9-11-26.xlsx"),
]
SOURCE_FILE = next((c for c in CANDIDATE_SOURCES if os.path.exists(c)), CANDIDATE_SOURCES[0])

# Candidate Reference Templates (.xlsm)
CANDIDATE_TEMPLATES = [
    r"C:\Users\vishal walunj\Downloads\Reference_Mapping_Files\CreateBillingEventsTemplate Paid-Unpaid Mapping.xlsm",
    r"C:\Users\vishal walunj\Vishal\1D_AIR_CONTROL_CONCEPTS\AiretechProjects\DEV1\DataFiles\BillingEvents Import\PAID CreateBillingEventsTemplate DEV1 V1.xlsm",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "CreateBillingEventsTemplate Paid-Unpaid Mapping.xlsm")
]
REFERENCE_TEMPLATE = next((c for c in CANDIDATE_TEMPLATES if os.path.exists(c)), None)

# Target Output Files (.xlsx Clean Data FBDI Workbooks)
OUT_PAID_XLSX = os.path.join(OUTPUT_DATA_DIR, "PaidBillingEvents_Populated.xlsx")
OUT_UNPAID_XLSX = os.path.join(OUTPUT_DATA_DIR, "UnpaidBillingEvents_Populated.xlsx")

# Target Output Files (.xlsm Macro-Enabled FBDI Workbooks)
OUT_PAID_XLSM = os.path.join(OUTPUT_DATA_DIR, "PaidCreateBillingEventsTemplate_FINAL_LOAD.xlsm")
OUT_UNPAID_XLSM = os.path.join(OUTPUT_DATA_DIR, "UnPaidCreateBillingEventsTemplate_FINAL_LOAD.xlsm")

# Target Output Files (.csv Interface Files)
OUT_PAID_CSV = os.path.join(OUTPUT_DATA_DIR, "Paid_PjbBillingEventsXface.csv")
OUT_UNPAID_CSV = os.path.join(OUTPUT_DATA_DIR, "Unpaid_PjbBillingEventsXface.csv")

# =============================================================================
# 2. DEFAULT CONSTANTS & HEADERS (78 Columns Oracle FBDI)
# =============================================================================
SOURCE_NAME = "Data Conversion"
ORGANIZATION = "Airetech"
CONTRACT_TYPE = "Airetech_Customer Contract for Project"
CONTRACT_LINE_NUM = "Conversion Contract Line"
EVENT_TYPE = "AIR Conversion"
CURRENCY_CODE = "USD"
CUTOVER_DATE = "07/31/2026"

BILLING_EVENT_HEADERS = [
    'Source Name', 'Source Reference', 'Organization', '*Contract Type', '*Contract Number',
    '*Contract Line Number', '*Event Type', 'Description', 'Completion Date',
    '*Bill Transaction Currency ', 'Event Amount in Bill Transaction Currency',
    'Project Number', 'Transaction Task Number', 'Hold Invoice', 'Hold Revenue',
    'Attribute Category', 'Custom Element 1', 'Custom Element 2', 'Custom Element 3',
    'Custom Element 4', 'Custom Element 5', 'Custom Element 6', 'Custom Element 7',
    'Custom Element 8', 'Custom Element 9', 'Custom Element 10', 'Custom Element 11',
    'Custom Element 12', 'Custom Element 13', 'Custom Element 14', 'Custom Element 15',
    'Custom Element 16', 'Custom Element 17', 'Custom Element 18', 'Custom Element 19',
    'Custom Element 20', 'Custom Element 21', 'Custom Element 22', 'Custom Element 23',
    'Custom Element 24', 'Custom Element 25', 'Custom Element 26', 'Custom Element 27',
    'Custom Element 28', 'Custom Element 29', 'Custom Element 30', 'Custom Number Element 1',
    'Custom Number Element 2', 'Custom Number Element 3', 'Custom Number Element 4',
    'Custom Number Element 5', 'Custom Number Element 6', 'Custom Number Element 7',
    'Custom Number Element 8', 'Custom Number Element 9', 'Custom Number Element 10',
    'Custom Date Element 1', 'Custom Date Element 2', 'Custom Date Element 3',
    'Custom Date Element 4', 'Custom Date Element 5', 'Custom Date Element 6',
    'Custom Date Element 7', 'Custom Date Element 8', 'Custom Date Element 9',
    'Custom Date Element 10', 'Custom Timestamp Element 1', 'Custom Timestamp Element 2',
    'Custom Timestamp Element 3', 'Custom Timestamp Element 4', 'Custom Timestamp Element 5',
    'Reverse in Next Period', 'Item Based Event', 'Quantity', 'Item Number',
    'Unit of measure', 'Unit Price', 'Prepayment Request Billing Event'
]


# =============================================================================
# 3. HELPER FUNCTIONS
# =============================================================================
def clean_job_no(val):
    """Preserves original job number string, padding numeric values to 10 digits."""
    if pd.isna(val):
        return ""
    val_str = str(val).strip()
    if val_str.endswith('.0'):
        val_str = val_str[:-2]
    if val_str.isdigit():
        return val_str.zfill(10)
    return val_str


def parse_numeric(val):
    """Parses numeric amounts safely handling commas, currency symbols, and strings."""
    if pd.isna(val):
        return 0.0
    val_str = str(val).replace(',', '').replace('$', '').strip()
    try:
        return float(val_str)
    except Exception:
        return 0.0


def safe_save(wb, path):
    """Saves workbook safely, creating a _clean alternative if file is locked in Excel."""
    try:
        wb.save(path)
        print(f"  [OK] Successfully saved: {path}")
        return path
    except Exception as e:
        base, ext = os.path.splitext(path)
        alt = f"{base}_clean{ext}"
        wb.save(alt)
        print(f"  [WARN] File locked in Excel ({e}): Saved to alternative -> {alt}")
        return alt


def create_standalone_billing_workbook(rows_list):
    """Generates a complete, cleanly styled Oracle Billing Events FBDI workbook (.xlsx)."""
    wb = openpyxl.Workbook()
    # Sheet 1: Instructions and CSV Generation
    ws_inst = wb.active
    ws_inst.title = "Instructions and CSV Generation"
    ws_inst.cell(row=1, column=1, value="Oracle Project Billing Events Interface")

    # Sheet 2: Project Billing Events
    ws_events = wb.create_sheet(title="Project Billing Events")
    ws_events.views.sheetView[0].showGridLines = True

    fill_hdr = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font_hdr = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_reg = Font(name="Segoe UI", size=9.5, color="000000")
    thin_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # Row 2 & 3
    ws_events.cell(row=2, column=1, value="Project Billing Events")
    ws_events.cell(row=3, column=1, value="* Required")

    # Row 4: Headers
    ws_events.row_dimensions[4].height = 24
    for c_idx, h in enumerate(BILLING_EVENT_HEADERS, start=1):
        cell = ws_events.cell(row=4, column=c_idx, value=h)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Rows 5+: Data Rows
    for r_idx, r_data in enumerate(rows_list, start=5):
        ws_events.row_dimensions[r_idx].height = 20
        col_vals = [
            r_data.get("Source Name"),
            r_data.get("Source Reference"),
            r_data.get("Organization"),
            r_data.get("Contract Type"),
            r_data.get("Contract Number"),
            r_data.get("Contract Line Number"),
            r_data.get("Event Type"),
            r_data.get("Description"),
            r_data.get("Completion Date"),
            r_data.get("Bill Transaction Currency"),
            r_data.get("Event Amount in Bill Transaction Currency"),
            r_data.get("Project Number")
        ]
        # Pad up to 78 columns
        col_vals.extend([None] * (len(BILLING_EVENT_HEADERS) - len(col_vals)))

        for c_idx, val in enumerate(col_vals, start=1):
            cell = ws_events.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_reg
            cell.border = thin_border
            if c_idx == 11 and val is not None:
                cell.number_format = "$#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif val is not None:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    return wb


def populate_xlsm_template(template_path, rows_list, output_path):
    """Populates the reference .xlsm macro-enabled template starting at Row 5."""
    if not template_path or not os.path.exists(template_path):
        return None
    try:
        wb = openpyxl.load_workbook(template_path, keep_vba=True)
        ws = wb["Project Billing Events"]
        ws.views.sheetView[0].showGridLines = True

        # Clear existing rows starting from Row 5
        max_r = max(ws.max_row, 20)
        max_c = len(BILLING_EVENT_HEADERS)
        for r in range(5, max_r + 1):
            for c in range(1, max_c + 1):
                ws.cell(row=r, column=c).value = None

        font_reg = Font(name="Segoe UI", size=9.5, color="000000")
        thin_side = Side(border_style="thin", color="D9D9D9")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        for r_idx, r_data in enumerate(rows_list, start=5):
            ws.row_dimensions[r_idx].height = 20
            col_vals = [
                r_data.get("Source Name"),
                r_data.get("Source Reference"),
                r_data.get("Organization"),
                r_data.get("Contract Type"),
                r_data.get("Contract Number"),
                r_data.get("Contract Line Number"),
                r_data.get("Event Type"),
                r_data.get("Description"),
                r_data.get("Completion Date"),
                r_data.get("Bill Transaction Currency"),
                r_data.get("Event Amount in Bill Transaction Currency"),
                r_data.get("Project Number")
            ]
            col_vals.extend([None] * (max_c - len(col_vals)))

            for c_idx, val in enumerate(col_vals, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = font_reg
                cell.border = thin_border
                if c_idx == 11 and val is not None:
                    cell.number_format = "$#,##0.00"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif val is not None:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        return safe_save(wb, output_path)
    except Exception as e:
        print(f"  [WARN] Could not populate .xlsm template ({e}). Skipping .xlsm output.")
        return None


def export_interface_csv(rows_list, output_csv_path):
    """Exports raw interface CSV matching Oracle PjbBillingEventsXface format."""
    lines = []
    for r in rows_list:
        fields = [
            str(r.get("Source Name", "")),
            str(r.get("Source Reference", "")),
            str(r.get("Organization", "")),
            str(r.get("Contract Type", "")),
            str(r.get("Contract Number", "")),
            str(r.get("Contract Line Number", "")),
            str(r.get("Event Type", "")),
            f'"{r.get("Description", "")}"' if "," in str(r.get("Description", "")) else str(r.get("Description", "")),
            str(r.get("Completion Date", "")),
            str(r.get("Bill Transaction Currency", "")),
            f'{r.get("Event Amount in Bill Transaction Currency", 0):.2f}',
            str(r.get("Project Number", ""))
        ]
        # Pad remaining columns up to 77, ending with END marker
        fields.extend([""] * (77 - len(fields)))
        fields.append("END")
        lines.append(",".join(fields))

    with open(output_csv_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print(f"  [OK] Successfully saved CSV: {output_csv_path}")


# =============================================================================
# 4. MAIN CONVERSION PROCESS
# =============================================================================
def main():
    print("=" * 80)
    print("  Airetech PPM Conversion - Billing Events FBDI Generation Starting")
    print("=" * 80)

    if not os.path.exists(SOURCE_FILE):
        raise FileNotFoundError(f"Source file not found at: {SOURCE_FILE}")

    # Load Source File
    df_src = pd.read_excel(SOURCE_FILE)
    df_src = df_src[df_src['Job No.'].notna() & (df_src['Job No.'].astype(str).str.strip() != 'Job No.')].reset_index(drop=True)
    print(f"  [OK] Loaded source dataset: {SOURCE_FILE} ({len(df_src)} records)")

    # Data Collections
    paid_event_rows = []
    unpaid_event_rows = []

    for i in range(len(df_src)):
        row = df_src.iloc[i]
        j_num = clean_job_no(row['Job No.'])
        j_name = str(row['Job Name']).strip() if ('Job Name' in row and pd.notna(row['Job Name'])) else f"Job_{j_num}"

        rev_billed = parse_numeric(row.get('Rev Billed', 0))
        ar_owing = parse_numeric(row.get('A/R Owing', 0))

        paid_amt = round(rev_billed - ar_owing, 2)
        unpaid_amt = round(ar_owing, 2)

        # 1. Paid Billing Event (Rev Billed - A/R Owing > 0)
        if paid_amt > 0:
            paid_event_rows.append({
                "Source Name": SOURCE_NAME,
                "Source Reference": f"EVT_PAID_{j_num}",
                "Organization": ORGANIZATION,
                "Contract Type": CONTRACT_TYPE,
                "Contract Number": j_num,
                "Contract Line Number": CONTRACT_LINE_NUM,
                "Event Type": EVENT_TYPE,
                "Description": f"Paid Invoice Event - Job {j_num} ({j_name})",
                "Completion Date": CUTOVER_DATE,
                "Bill Transaction Currency": CURRENCY_CODE,
                "Event Amount in Bill Transaction Currency": paid_amt,
                "Project Number": j_num
            })

        # 2. Unpaid Billing Event (A/R Owing > 0)
        if unpaid_amt > 0:
            unpaid_event_rows.append({
                "Source Name": SOURCE_NAME,
                "Source Reference": f"EVT_UNPAID_{j_num}",
                "Organization": ORGANIZATION,
                "Contract Type": CONTRACT_TYPE,
                "Contract Number": j_num,
                "Contract Line Number": CONTRACT_LINE_NUM,
                "Event Type": EVENT_TYPE,
                "Description": f"Unpaid Invoice Event - Job {j_num} ({j_name})",
                "Completion Date": CUTOVER_DATE,
                "Bill Transaction Currency": CURRENCY_CODE,
                "Event Amount in Bill Transaction Currency": unpaid_amt,
                "Project Number": j_num
            })

    tot_paid = sum(r["Event Amount in Bill Transaction Currency"] for r in paid_event_rows)
    tot_unpaid = sum(r["Event Amount in Bill Transaction Currency"] for r in unpaid_event_rows)

    print(f"\n  [OK] Paid Billing Events:   {len(paid_event_rows)} Records (Total: ${tot_paid:,.2f})")
    print(f"  [OK] Unpaid Billing Events: {len(unpaid_event_rows)} Records (Total: ${tot_unpaid:,.2f})")

    # -------------------------------------------------------------------------
    # 1. Save Standalone Populated Excel Files (.xlsx)
    # -------------------------------------------------------------------------
    print("\n  [1/3] Generating Separate Paid and Unpaid FBDI Workbooks (.xlsx)...")
    wb_paid = create_standalone_billing_workbook(paid_event_rows)
    safe_save(wb_paid, OUT_PAID_XLSX)

    wb_unpaid = create_standalone_billing_workbook(unpaid_event_rows)
    safe_save(wb_unpaid, OUT_UNPAID_XLSX)

    # -------------------------------------------------------------------------
    # 2. Populate Macro-Enabled Templates (.xlsm) if base template available
    # -------------------------------------------------------------------------
    if REFERENCE_TEMPLATE and os.path.exists(REFERENCE_TEMPLATE):
        print("\n  [2/3] Generating Separate Paid and Unpaid Macro-Enabled Files (.xlsm)...")
        populate_xlsm_template(REFERENCE_TEMPLATE, paid_event_rows, OUT_PAID_XLSM)
        populate_xlsm_template(REFERENCE_TEMPLATE, unpaid_event_rows, OUT_UNPAID_XLSM)
    else:
        print("\n  [2/3] No base .xlsm template found. Standalone .xlsx workbooks generated.")

    # -------------------------------------------------------------------------
    # 3. Export Interface CSV Files
    # -------------------------------------------------------------------------
    print("\n  [3/3] Exporting Interface CSV Files...")
    export_interface_csv(paid_event_rows, OUT_PAID_CSV)
    export_interface_csv(unpaid_event_rows, OUT_UNPAID_CSV)

    print("\n" + "=" * 80)
    print("  [OK] BILLING EVENTS CONVERSION COMPLETED SUCCESSFULLY!")
    print(f"  Target Directory: {OUTPUT_DATA_DIR}")
    print(f"  Paid Events File:   {OUT_PAID_XLSX} ({len(paid_event_rows)} records, ${tot_paid:,.2f})")
    print(f"  Unpaid Events File: {OUT_UNPAID_XLSX} ({len(unpaid_event_rows)} records, ${tot_unpaid:,.2f})")
    print("=" * 80)


if __name__ == "__main__":
    main()